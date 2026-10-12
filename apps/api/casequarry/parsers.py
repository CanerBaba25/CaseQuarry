import csv
import io
import json
import mimetypes
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass, field
from datetime import UTC
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path

from bs4 import BeautifulSoup


@dataclass
class Parsed:
    item_type: str
    title: str
    text: str = ""
    state: str = "indexed"
    error: str | None = None
    pages: list[tuple[int, int, int]] = field(default_factory=list)
    date_raw: str = ""
    date_utc: object = None
    timezone_source: str = "missing"
    sender: str = ""
    recipients: str = ""
    email: dict = field(default_factory=dict)
    metadata: dict = field(default_factory=dict)
    attachments: list[tuple[str, bytes]] = field(default_factory=list)


def safe_text_html(value: str) -> str:
    document = BeautifulSoup(value, "html.parser")
    for node in document(["script", "style", "iframe", "object", "embed", "form", "svg", "link", "meta"]):
        node.decompose()
    return document.get_text("\n", strip=True)


def addresses(value: str) -> list[dict]:
    return [{"name": name.strip(), "email": address.strip().lower()} for name, address in getaddresses([value]) if address]


def parse_email(content: bytes, filename: str, max_chars: int) -> Parsed:
    message = BytesParser(policy=policy.default).parsebytes(content)
    result = Parsed("email", str(message.get("Subject", "(No subject)")))
    result.email = {"from": addresses(str(message.get("From", ""))), "to": addresses(str(message.get("To", ""))), "cc": addresses(str(message.get("Cc", ""))), "bcc": addresses(str(message.get("Bcc", ""))), "reply_to": addresses(str(message.get("Reply-To", ""))), "raw_headers": "\n".join(f"{key}: {value}" for key, value in message.raw_items()), "message_id": str(message.get("Message-ID", "")), "in_reply_to": str(message.get("In-Reply-To", "")), "references": re.findall(r"<[^>]+>", str(message.get("References", ""))), "sent_at_raw": str(message.get("Date", ""))}
    result.sender = ", ".join(item["email"] for item in result.email["from"])
    result.recipients = ", ".join(item["email"] for item in result.email["to"] + result.email["cc"] + result.email["bcc"])
    result.date_raw = result.email["sent_at_raw"]
    try:
        parsed = parsedate_to_datetime(result.date_raw)
        if parsed.tzinfo:
            result.date_utc, result.timezone_source = parsed.astimezone(UTC), "explicit"
        else:
            result.timezone_source = "unknown"
            result.metadata["date_warning"] = "Timestamp has no explicit timezone; UTC was not inferred."
    except (TypeError, ValueError, OverflowError):
        if result.date_raw:
            result.metadata["date_warning"] = "The original date could not be normalized."
    body = message.get_body(preferencelist=("plain", "html")) if message.is_multipart() else message
    if body and body.get_content_type() in {"text/plain", "text/html"}:
        decoded = body.get_content()
        result.text = safe_text_html(decoded) if body.get_content_type() == "text/html" else str(decoded)
    for index, part in enumerate(message.iter_attachments() if message.is_multipart() else []):
        if index >= 200:
            result.state, result.error = "partial", "ATTACHMENT_LIMIT"
            break
        name = str(part.get_filename() or f"attachment-{index + 1}")
        if part.get_content_type() == "message/rfc822":
            payload = part.get_payload()
            data = payload[0].as_bytes(policy=policy.default) if isinstance(payload, list) and payload else b""
            if not name.lower().endswith(".eml"):
                name += ".eml"
        else:
            data = part.get_payload(decode=True) or b""
        result.attachments.append((Path(name.replace("\\", "/")).name[:400], data))
    if len(result.text) > max_chars:
        result.text, result.state, result.error = result.text[:max_chars], "partial", "TEXT_LIMIT"
    if message.defects:
        result.metadata["mime_warnings"] = [type(defect).__name__ for defect in message.defects]
    return result


def parse_document(path: Path, filename: str, settings) -> Parsed:
    extension = Path(filename).suffix.lower()
    media_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    try:
        if extension in {".docx", ".xlsx", ".pptx"}:
            from .ingestion import safe_zip_entries
            with zipfile.ZipFile(path) as archive:
                safe_zip_entries(archive, settings)
        if extension == ".eml":
            return parse_email(path.read_bytes(), filename, settings.max_extracted_chars)
        if extension in {".mbox", ".zip"}:
            return Parsed("container", filename, state="stored", metadata={"format": extension[1:]})
        if extension == ".pdf":
            from pypdf import PdfReader
            with path.open("rb") as header:
                is_pdf = header.read(5) == b"%PDF-"
            if not is_pdf:
                return Parsed("pdf", filename, state="failed", error="MIME_MISMATCH")
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                return Parsed("pdf", filename, state="locked", error="PASSWORD_REQUIRED")
            result = Parsed("pdf", filename, metadata={"page_count": len(reader.pages)})
            parts, offset = [], 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if offset + len(text) > settings.max_extracted_chars:
                    result.state, result.error = "partial", "TEXT_LIMIT"
                    break
                result.pages.append((number, offset, offset + len(text)))
                parts.append(text)
                offset += len(text) + 2
            result.text = "\n\n".join(parts)
            if not result.text.strip():
                result.state, result.error = "partial", "OCR_REQUIRED"
            return result
        if extension == ".docx":
            from docx import Document
            document = Document(path)
            text = "\n".join([paragraph.text for paragraph in document.paragraphs] + ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
            result = Parsed("document", filename, text=text)
        elif extension == ".xlsx":
            return parse_spreadsheet(path, filename, settings)
        elif extension == ".pptx":
            from pptx import Presentation
            slides = Presentation(path).slides
            result = Parsed("presentation", filename, text="\n\n".join(f"Slide {number}\n" + "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")) for number, slide in enumerate(slides, 1)))
        elif extension in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
            from PIL import Image
            with Image.open(path) as image:
                image.verify()
            if shutil.which("tesseract"):
                command = subprocess.run(["tesseract", str(path), "stdout"], capture_output=True, timeout=60, check=False)
                if command.returncode == 0:
                    result = Parsed("image", filename, text=command.stdout.decode("utf-8", errors="replace"), metadata={"ocr_engine": "tesseract", "ocr_quality": "unverified"})
                else:
                    result = Parsed("image", filename, state="partial", error="OCR_FAILED")
            else:
                result = Parsed("image", filename, state="partial", error="OCR_DEPENDENCY_MISSING")
        elif extension in {".txt", ".md", ".html", ".htm", ".csv", ".json", ".log"}:
            with path.open("rb") as stream:
                data = stream.read(settings.max_extracted_chars * 4 + 1)
            text = data.decode("utf-8-sig", errors="replace")
            if extension in {".html", ".htm"}:
                text = safe_text_html(text)
            elif extension == ".json":
                json.loads(text)
            elif extension == ".csv":
                text = "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text)))
            result = Parsed("text", filename, text=text)
            if "\ufffd" in text:
                result.state, result.error = "partial", "ENCODING_REPLACEMENTS"
        else:
            return Parsed("file", filename, state="unsupported", error="PARSER_UNAVAILABLE", metadata={"media_type": media_type, "format": extension})
        if len(result.text) > settings.max_extracted_chars:
            result.text, result.state, result.error = result.text[:settings.max_extracted_chars], "partial", "TEXT_LIMIT"
        if not result.text.strip() and result.state == "indexed":
            result.state, result.error = "partial", "NO_EXTRACTED_TEXT"
        return result
    except Exception as error:
        return Parsed("file" if extension not in {".pdf", ".eml"} else extension[1:], filename, state="failed", error=f"PARSE_FAILED:{type(error).__name__}")


def parse_spreadsheet(path: Path, filename: str, settings) -> Parsed:
    from openpyxl import load_workbook

    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            parts, length = [], 0
            for sheet in workbook:
                parts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join("" if value is None else str(value) for value in row)
                    length += len(line)
                    if length > settings.max_extracted_chars:
                        return Parsed("spreadsheet", filename, "\n".join(parts), "partial", "TEXT_LIMIT")
                    parts.append(line)
            return Parsed("spreadsheet", filename, text="\n".join(parts))
        finally:
            workbook.close()
